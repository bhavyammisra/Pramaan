"""v6 Pramaan screening deck (v5 + review fixes, D-011): plain, text-led, fewer boxes, no icon badges, no invented scores.
Usage: python scripts/build_deck_v6.py <BIOS_V2_CORE_FRONTIER_Screening_Template.pptx> <out.pptx>
Fixes from the v4 review: verified numbers only, consistent worked example (71x usual), honest Pramaan claims,
template wording on slide 2, a simple team-drawn diagram on slide 4, 48-hour plan on slide 5."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from build_deck import (NAVY, ICE, PANEL, TEXT, MUTED, GREEN, AMBER, RED, WHITE, HEAD, rgb,
                        add_par, textbox, box, arrow, strip_instructions, heading, sources, notes)

L = PP_ALIGN.LEFT
MONO = "Consolas"
LINE = "C9CEDB"


def text(s, x, y, w, h, name, pars, anchor=MSO_ANCHOR.TOP):
    """pars: list of (runs, space_after) paragraphs; runs as in build_deck.add_par."""
    tb, tf = textbox(s, x, y, w, h, name, anchor=anchor)
    for i, (runs, after) in enumerate(pars):
        add_par(tf, runs, align=L, first=i == 0, space_after=after)
    return tf


def rule(s, x, y, w, h, name, color=NAVY):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = rgb(color); r.line.fill.background(); r.shadow.inherit = False
    r.name = name
    return r


def cap(t, color=NAVY):
    return [(t, 10.5, True, color)]


# ---------------------------------------------------------------- slide 1
def slide1(s):
    heading(s, headline="> Fraud has moved past OTPs: victims now approve the payment themselves")
    _, tf = box(s, 4.15, 0.66, 8.7, 0.9, "Team strip", fill=PANEL, radius=0.04, pad=0.12)
    add_par(tf, [("Team Caffeinated Nerds", 11.5, True, NAVY), ("   ·   Challenge: ", 11, False, NAVY), ("Frontier · PS-04", 11, True, NAVY),
                 (" · Fast and Secure Online Transaction Authentication", 11, False, NAVY)], first=True, align=L)
    for i, (name, roll) in enumerate([("Bhavyam Misra", "1024030535"), ("Gunagye Jain", "1024030019"),
                                      ("Abhijay Nagal", "1024170013"), ("Shefali Sharma", "1024030284")]):
        tb, ttf = textbox(s, 4.27 + i * 2.13, 1.01, 2.05, 0.45, f"Member {i+1}")
        add_par(ttf, [(name, 10.5, True, TEXT)], first=True, align=L)
        add_par(ttf, [("Roll No: " + roll, 10, False, MUTED)], align=L)

    text(s, 0.5, 2.8, 6.6, 1.25, "Problem", [
        (cap("THE PROBLEM"), 3),
        ([("Asha, 72, gets a call from a fake “CBI officer”. On her own phone, at home, she types her own UPI PIN. "
           "Right person, right phone, right PIN: the bank ", 13, False, TEXT), ("approves ₹85,000 to a scammer", 13, True, RED), (".", 13, False, TEXT)], 0)])
    rule(s, 0.5, 3.88, 0.06, 0.62, "Insight bar")
    text(s, 0.72, 3.88, 6.4, 0.62, "Insight", [([("A PIN checks ", 18, False, NAVY, False, HEAD), ("who", 18, True, NAVY, False, HEAD),
                                                 (" you are. Nobody checks ", 18, False, NAVY, False, HEAD), ("what's happening", 18, True, NAVY, False, HEAD),
                                                 (".", 18, False, NAVY, False, HEAD)], 0)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.5, 4.68, 6.6, 1.7, "Stakeholder and why", [
        (cap("WHO IT'S FOR"), 2),
        ([("Banks, UPI apps and payment gateways", 12.5, True, TEXT), (", who must answer well inside UPI's 15-second response limit, on behalf of everyday payers, especially seniors and first-time users.", 12.5, False, TEXT)], 8),
        (cap("WHY IT'S WORTH SOLVING"), 2),
        ([("RBI's proposed fix is a 1-hour delay on every payment above ₹10,000. That slows every honest user too. A check that reads the situation can stop the scam and leave honest payments alone.", 12.5, False, TEXT)], 0)])

    text(s, 7.55, 2.8, 5.3, 0.3, "Evidence label", [(cap("EVIDENCE"), 0)])
    stats = [("₹22,495 cr", "lost to cyber fraud in India in 2025 (MHA, NCRP data)"),
             ("85%", "of that came from investment (76%) and “digital arrest” (9%) scams, where victims send the money themselves (MHA)"),
             ("98.5%", "of fraud value is in payments above ₹10,000; account-takeover fraud is “now negligible” (RBI, 9 Apr 2026)")]
    for i, (big, small) in enumerate(stats):
        y = 3.12 + i * 1.08
        rule(s, 7.55, y + 0.05, 0.05, 0.9, f"Stat {i+1} bar", color=ICE)
        text(s, 7.75, y, 5.1, 1.0, f"Stat {i+1}", [([(big, 24, True, NAVY, False, HEAD)], 0), ([(small, 11.5, False, TEXT)], 0)])

    sdgs = [("00689D", "PRIMARY · SDG 16", "16.4: cut illicit financial flows & organised crime"),
            ("A21942", "SUPPORTING · SDG 8", "8.10: access to financial services for all"),
            ("DD5A1F", "SUPPORTING · SDG 9", "9.1: reliable, resilient (digital) infrastructure")]
    for i, (col, head, body) in enumerate(sdgs):
        _, tf = box(s, 0.5 + i * 4.175, 6.5, 4.0, 0.56, f"SDG {i+1}", fill=col, radius=0.03, pad=0.12, anchor=MSO_ANCHOR.MIDDLE)
        add_par(tf, [(head, 10, True, WHITE)], first=True, align=L)
        add_par(tf, [(body, 10, False, WHITE)], align=L)
    sources(s, "Sources: MHA/NCRP via ThePrint (2025 losses by scam type); RBI Discussion Paper “Exploring safeguards in digital payments to curb frauds”, 9 Apr 2026. "
               "Links: github.com/bhavyammisra/BIOS-Transaction-Authentication")
    notes(s, "We're team Caffeinated Nerds, on Frontier problem PS-04: Fast and Secure Online Transaction Authentication. "
             "Meet Asha. She's 72. A fake CBI officer calls her. On her own phone, at home, she types her own UPI PIN, and the bank approves 85,000 rupees to a scammer. "
             "Every check passed: right person, right phone, right PIN. A PIN checks who you are. Nobody checks what's happening. "
             "This is now the main problem. Indians lost about 22,500 crore rupees to cyber fraud in 2025, and 85 percent of that came from investment and digital-arrest scams, where the victim sends the money. "
             "The RBI says account-takeover fraud is now negligible, and 98.5 percent of fraud value is in payments above 10,000 rupees. "
             "We build for banks, UPI apps and gateways, who must answer well inside UPI's 15-second response limit. RBI's proposed fix delays every payment above 10,000 rupees by an hour, which slows every honest user. We want precision instead. "
             "Primary SDG 16.4, supported by 8.10 and 9.1.")


# ---------------------------------------------------------------- slide 2
def slide2(s):
    heading(s, headline="> Today's defences check the credential or the receiver, never the moment")
    cards = [
        ("Static two-factor auth", "UPI PIN · SMS OTP · 3-D Secure",
         "Device binding plus a PIN or OTP on every payment. RBI keeps 2FA as the minimum from April 2026.",
         "Checks the credential, not the context. A coached victim types it in; a SIM swap steals it. A ₹50 tea and a ₹2 lakh transfer get the same check."),
        ("Commercial risk engines", "Stripe Radar · Sift · BioCatch",
         "Machine-learning scores over hundreds of signals; the merchant or bank sets block, review or step-up rules.",
         "Closed, inside one provider's network. The output is a score for an analyst, not a reason for the customer. A high score triggers the same OTP or 3DS whatever the kind of risk."),
        ("India's shared intelligence", "NPCI FRM · DoT FRI · RBI MuleHunter.AI",
         "Flags risky phone numbers and mule accounts across banks and UPI apps.",
         "Works after the fact: a number must be reported first, and scam numbers live for days. It scores the receiver, not what is happening in this session."),
    ]
    for i, (title, sub, does, gap) in enumerate(cards):
        x = 0.5 + i * 4.19
        rule(s, x, 2.82, 3.97, 0.04, f"Approach {i+1} rule")
        text(s, x, 2.95, 3.97, 2.95, f"Approach {i+1}", [
            ([(f"{i+1}. {title}", 15, True, NAVY, False, HEAD)], 0),
            ([(sub, 10.5, False, MUTED, True)], 8),
            (cap("What it does"), 1),
            ([(does, 13, False, TEXT)], 10),
            (cap("What's still missing", RED), 1),
            ([(gap, 13, False, TEXT)], 0)])
    _, tf = box(s, 0.5, 6.0, 12.35, 0.98, "Gap statement", fill=NAVY, radius=0.03, pad=0.2, anchor=MSO_ANCHOR.MIDDLE)
    add_par(tf, [("The gap we are addressing is: ", 14, True, ICE, False, HEAD),
                 ("no lightweight, real-time service combines behaviour, device and transaction context into an explained allow / challenge / block decision, "
                  "and picks the kind of challenge that actually stops the risk in front of it.", 13, False, WHITE)], first=True, align=L)
    sources(s, "Sources: RBI Authentication Directions (25 Sep 2025) & Discussion Paper (9 Apr 2026); Stripe Radar risk-evaluation docs; Sift; BioCatch; NPCI FRM; DoT FRI (May 2025). "
               "Full links: repo › docs/02-research/prior-art.md")
    notes(s, "Three families of solutions exist today. "
             "First, static two-factor auth: PIN, OTP, 3-D Secure. It's mandatory, but it checks the credential, not the context. A coached victim simply types it in. "
             "Second, commercial risk engines like Stripe Radar, Sift and BioCatch. They're strong, but closed, the output is a score for an analyst, and a high score triggers the same OTP whatever the kind of risk. "
             "Third, India's shared intelligence: NPCI's fraud monitoring, DoT's Fraud Risk Indicator, RBI's MuleHunter. They flag risky numbers and accounts, but only after someone reports them, and they score the receiver, not this session. "
             "So the gap: nothing lightweight combines all three signal families into an explained decision and picks the challenge that actually stops the risk.")


# ---------------------------------------------------------------- slide 3
def slide3(s):
    heading(s, headline="> Pramaan: a PIN checks who you are. We check what's happening, before every payment.")
    text(s, 0.5, 2.8, 5.85, 3.0, "Solution", [
        (cap("WHAT WE BUILD"), 2),
        ([("A cloud risk API that a bank or UPI app calls ", 13, False, TEXT), ("before", 13, True, TEXT),
          (" the PIN step. It returns allow, challenge or block in milliseconds, with reasons in plain language. It sits on top of mandated 2FA and never replaces it.", 13, False, TEXT)], 10),
        (cap("WHAT'S DIFFERENT"), 2),
        ([("1  ", 12.5, True, NAVY), ("Each payment is compared with ", 12.5, False, TEXT), ("this user's own", 12.5, True, TEXT), (" history, not a global average.", 12.5, False, TEXT)], 4),
        ([("2  ", 12.5, True, NAVY), ("Amount-aware: small payments are almost never interrupted; a large first-time transfer is checked closely.", 12.5, False, TEXT)], 4),
        ([("3  ", 12.5, True, NAVY), ("The challenge fits the risk. A passkey on the owner's registered phone stops an impostor. Only a pause and a second person stop a coached victim.", 12.5, False, TEXT)], 0)])
    text(s, 0.5, 5.62, 5.85, 0.3, "Asha label", [(cap("WHAT ASHA SEES IN SITUATION 3"), 0)])
    rule(s, 0.5, 5.97, 0.05, 0.78, "Quote bar", color=AMBER)
    text(s, 0.7, 5.95, 5.65, 1.12, "Quote", [
        ([("“Police, CBI or RBI never ask you to move money over a call. We've paused this payment for 30 minutes and asked Rohan to confirm. You can cancel anytime.”", 12, False, TEXT, True)], 3),
        ([("Rohan, her son, gets Approve / Cancel on Telegram. The payment goes only if he approves; no reply in 30 minutes cancels it.", 10.5, False, MUTED)], 0)])

    X0, W = 6.65, 6.2
    text(s, X0, 2.8, W, 0.3, "Scen label", [(cap("THE SAME ₹85,000 PAYMENT, FOUR SITUATIONS"), 0)])
    cols = [(X0, 2.95, "WHAT WE SEE"), (X0 + 3.05, 1.2, "LIKELY RISK"), (X0 + 4.3, 1.9, "RESPONSE")]
    for x, w, h in cols:
        text(s, x, 3.12, w, 0.25, f"Head {h}", [([(h, 9.5, True, MUTED)], 0)])
    rows = [("Usual payee, own phone, normal typing", "None", "Allow, no extra step", GREEN),
            ("New phone, SIM changed yesterday, new city", "Impostor", "Approve on registered phone (passkey)", AMBER),
            ("Own phone at home, on a call, new UPI ID pasted", "Coached victim", "30-min hold until family approves", AMBER),
            ("Emulator + VPN, 15 PIN retries in 10 min", "Bot", "Block", RED)]
    for i, (see, risk, resp, col) in enumerate(rows):
        y = 3.42 + i * 0.6
        rule(s, X0, y, W, 0.012, f"Row {i+1} line", color=LINE)
        text(s, X0, y + 0.05, 2.95, 0.52, f"Row {i+1} see", [([(see, 12, False, TEXT)], 0)], anchor=MSO_ANCHOR.MIDDLE)
        text(s, X0 + 3.05, y + 0.05, 1.2, 0.52, f"Row {i+1} risk", [([(risk, 12, False, TEXT)], 0)], anchor=MSO_ANCHOR.MIDDLE)
        rule(s, X0 + 4.3, y + 0.24, 0.12, 0.12, f"Row {i+1} mark", color=col)
        text(s, X0 + 4.5, y + 0.05, 1.7, 0.52, f"Row {i+1} resp", [([(resp, 12, True, col)], 0)], anchor=MSO_ANCHOR.MIDDLE)
    rule(s, X0, 3.42 + 4 * 0.6, W, 0.012, "Row end line", color=LINE)

    text(s, X0, 6.0, W, 1.15, "Reasons", [
        (cap("REASONS THE API RETURNS FOR SITUATION 3 (TOP 3)"), 3),
        ([("You're on a phone call while sending money", 11.5, False, TEXT)], 1),
        ([("First payment to this receiver", 11.5, False, TEXT)], 1),
        ([("Amount is 71× your usual ₹1,200", 11.5, False, TEXT)], 0)])
    notes(s, "Our solution is Pramaan, which means proof. It's a cloud risk API that a bank or UPI app calls just before the PIN step. "
             "It returns allow, challenge or block in milliseconds, with reasons in plain language, and it sits on top of mandated 2FA. "
             "Three things are different. We compare each payment with this user's own history. Thresholds depend on the amount, so small payments are almost never interrupted. "
             "And the challenge fits the risk. Take the same 85,000-rupee payment in four situations. Usual payee, own phone: allowed. "
             "New phone and a SIM changed yesterday: likely an impostor, so we ask for approval with a passkey on the owner's registered phone, which the new phone doesn't have. "
             "Own phone at home, but on a call and pasting a new UPI ID: likely a coached victim. An OTP won't help, because she'd type it in. So we pause the payment for 30 minutes with a scam warning, and her son Rohan gets Approve or Cancel on Telegram. It only goes through if he approves. "
             "An emulator retrying the PIN 15 times in 10 minutes: blocked. "
             "Each decision carries its reasons: here, she's on a call, it's a new receiver, and it's 71 times her usual amount.")


# ---------------------------------------------------------------- slide 4
def diagram_box(s, x, y, w, h, name, title, body, mono=False):
    b, tf = box(s, x, y, w, h, name, fill=WHITE, line=NAVY, radius=0.02, pad=0.11)
    add_par(tf, [(title, 10.5 if mono else 11.5, True, NAVY, False, MONO if mono else "Calibri")], first=True, align=L, space_after=3)
    for t in body:
        add_par(tf, [(t, 10.5, False, TEXT)], align=L, space_after=1)
    return b


def slide4(s):
    heading(s, headline="> One API call: context in, an explained decision out, in milliseconds")
    top, H = 2.8, 1.5
    diagram_box(s, 0.5, top, 2.05, H, "D1", "Demo UPI app + JS SDK",
                ["amount, payee, device", "typing rhythm, paste, time to pay", "on-call, screen-share (toggles)"])
    diagram_box(s, 2.85, top, 2.0, H, "D2", "POST /v1/authenticate",
                ["FastAPI", "error or > 150 ms: fall back to standard 2FA"], mono=True)
    diagram_box(s, 5.15, top, 2.0, H, "D3", "Feature builder",
                ["~25 features vs this user's own history", "e.g. amount ÷ usual, new payee, new device"])
    diagram_box(s, 7.45, top, 3.05, H, "D4", "Risk engine",
                ["1  rules: hard blocks, fast-path allow", "2  LightGBM: fraud probability p", "3  TreeSHAP: each signal's share → reasons",
                 "4  risk type: impostor, scam, bot, mule", "5  cost policy: decision + challenge type"])
    diagram_box(s, 10.8, top, 2.05, H, "D5", "Response",
                ["allow / challenge / block", "challenge type", "top-3 reasons", "latency_ms"])
    mid = top + 0.62
    for i, (x1, x2) in enumerate(((2.55, 2.85), (4.85, 5.15), (7.15, 7.45), (10.5, 10.8))):
        arrow(s, x1, mid, x2, mid, f"A{i+1}")
    # profile store under the feature builder
    diagram_box(s, 5.15, 4.6, 2.0, 0.78, "D3b", "Profile store", ["usual amount, devices, payees, hours"])
    arrow(s, 6.15, 4.6, 6.15, top + H, "A profile", width=1.25)
    # what happens after the decision (off the scoring path)
    b, tf = box(s, 7.45, 4.6, 5.4, 0.8, "After decision", fill=WHITE, line=MUTED, radius=0.02, pad=0.11)
    b.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_par(tf, [("After the decision, off the scoring path", 10.5, True, MUTED)], first=True, align=L, space_after=2)
    add_par(tf, [("fingerprint: ", 10.5, True, TEXT), ("WebAuthn passkey on the registered phone", 10.5, False, TEXT)], align=L)
    add_par(tf, [("family: ", 10.5, True, TEXT), ("n8n → Telegram Approve / Cancel → no reply in 30 min = cancel → label", 10.5, False, TEXT)], align=L)
    arrow(s, 11.82, top + H, 11.82, 4.6, "A after", width=1.25)
    text(s, 0.5, 4.6, 4.35, 0.85, "Built vs mocked", [
        ([("Built in the hack: ", 10.5, True, GREEN), ("everything in this diagram, plus a Supabase decision log and a React dashboard.", 10.5, False, TEXT)], 3),
        ([("Mocked: ", 10.5, True, AMBER), ("bank/NPCI rails, DoT FRI and mule feeds, call and screen-share state (a browser can't read them; a mobile SDK would).", 10.5, False, TEXT)], 0)])

    panels = [("STACK · ₹0 RUNNING COST",
               "Python 3.11, FastAPI, LightGBM, shap on Render free tier · React (Vite) on Vercel · Supabase · n8n self-hosted + Telegram Bot API · WebAuthn"),
              ("DATA",
               "No public dataset has behaviour, device and transaction signals together. Synthetic generator: 5,000 users, 30 days, 4 labelled attack types, "
               "plus honest look-alikes (travel, new phone, rent). Sanity check on IEEE-CIS and PaySim."),
              ("DECISION RULE",
               "Block if p ≥ 0.90 or a hard rule fires. Challenge if p ≥ ₹500 ÷ amount, kept between 0.05 and 0.60 (×0.6 for users 70+). "
               "Challenge type: lowest friction + p × amount × (1 − how well it stops this risk type).")]
    for i, (h, body) in enumerate(panels):
        x = 0.5 + i * 4.19
        rule(s, x, 5.7, 3.97, 0.03, f"Panel {i+1} rule")
        text(s, x, 5.8, 3.97, 1.4, f"Panel {i+1}", [(cap(h), 2), ([(body, 11, False, TEXT)], 0)])
    notes(s, "Here's how it works. The demo payment app, with our small JavaScript SDK, sends the payment, device info and session behaviour: typing rhythm, a pasted payee, time to pay. "
             "Call and screen-share state need a mobile SDK, so in the web demo they're toggles. "
             "Our FastAPI service validates the request. If it ever errors or takes more than 150 milliseconds, it falls back to the bank's standard 2FA, so it's never weaker than today's check. "
             "The feature builder computes about 25 features against this user's own history from the profile store. "
             "The risk engine runs hard rules first, then LightGBM gives a fraud probability, TreeSHAP gives each signal's share, which becomes the reasons and the risk type, and a cost policy picks the decision and the challenge. "
             "After the decision, off the scoring path: an impostor challenge is a WebAuthn passkey on the registered phone, and a family challenge goes through an n8n workflow to Telegram. "
             "The stack is free. The data is synthetic, and we say so: no public dataset has these signals together. "
             "The thresholds get stricter as the amount grows, and blocking needs 90 percent confidence or a hard rule.")


# ---------------------------------------------------------------- slide 5
def slide5(s):
    heading(s, headline="> A deployed, working prototype in the hack, honest about its limits")
    heads = ["1 · WHAT WE WILL BUILD", "2 · RISKS & LIMITATIONS", "3 · EXPECTED IMPACT"]
    for i, h in enumerate(heads):
        x = 0.5 + i * 4.19
        text(s, x, 2.8, 3.97, 0.32, f"Col head {i+1}", [([(h, 12.5, True, NAVY)], 0)])
        rule(s, x, 3.14, 3.97, 0.035, f"Col rule {i+1}")
    build = ["Live /v1/authenticate on Render, with the 2FA fail-safe",
             "Risk engine with reasons, trained on synthetic data",
             "Demo UPI app: registered-phone passkey and the n8n + Telegram family loop",
             "Dashboard: live decisions, their reasons, latency",
             "Replay: our policy vs RBI's blanket delay on the same traffic"]
    pars = [([(f"{k+1}  ", 12, True, NAVY), (t, 12, False, TEXT)], 5) for k, t in enumerate(build)]
    pars += [([("", 6, False, TEXT)], 0), (cap("48-HOUR PLAN"), 3)]
    pars += [([(h, 10.5, True, TEXT), ("  " + t, 10.5, False, TEXT)], 1) for h, t in
             (("H0–8", "data generator, API skeleton"), ("H8–20", "engine + /v1/authenticate"), ("H20–34", "app, passkeys, n8n loop"),
              ("H34–44", "dashboard, replay, load test"), ("H44–48", "deploy, rehearse"))]
    text(s, 0.5, 3.3, 3.97, 3.9, "Build body", pars)

    risks = [("Synthetic data", "The model learns our generator, not real fraud.", "Honest look-alikes, results per attack type, checks on public data; a bank pilot retrains on real history."),
             ("Call state needs an SDK", "A browser can't see calls or screen-sharing.", "A missing signal counts as “unknown”, never “safe”."),
             ("False positives", "Holding a genuine urgent payment hurts.", "Block only at p ≥ 0.90; challenge at most 5% of honest payments; cancel anytime."),
             ("Privacy (DPDP Act)", "Behaviour data is personal data.", "Derived features only, hashed device IDs, a retention window.")]
    pars = []
    for k, (h, prob, fix) in enumerate(risks):
        pars.append(([(h, 12, True, RED)], 0))
        pars.append(([(prob + " ", 11, False, TEXT), (fix, 11, False, MUTED, True)], 7 if k < 3 else 0))
    text(s, 4.69, 3.3, 3.97, 3.9, "Risks body", pars)

    mets = [("≥ 80%", "of scam-type fraud stopped, while challenging ≤ 5% of honest payments"),
            ("≥ 90%", "of honest ₹10,000+ transfers go through with no delay (a blanket rule delays all of them)"),
            ("< 50 ms", "p95 decision time, a small slice of UPI's 15-second response limit")]
    pars = [([("Targets for the prototype:", 11.5, True, TEXT)], 6)]
    for big, small in mets:
        pars.append(([(big, 22, True, NAVY, False, HEAD)], 0))
        pars.append(([(small, 11, False, TEXT)], 7))
    pars.append(([("Targets, measured on held-out synthetic data and shown live; the amount threshold is tuned so the 5% and 90% targets hold together. Baselines: OTP-only, a fixed threshold, RBI's blanket delay.", 10, False, MUTED, True)], 0))
    text(s, 8.88, 3.3, 3.97, 3.9, "Impact body", pars)
    notes(s, "In the hack we'll build a live authenticate API with the 2FA fail-safe, the risk engine with reasons, a demo UPI app with a real passkey check on the registered phone and the n8n and Telegram family loop, a dashboard, "
             "and a replay that runs our policy and RBI's blanket delay on the same traffic. "
             "The risks are real. Our data is synthetic, so we add honest look-alikes, report results per attack type and check against public data. "
             "A browser can't see call state, so a missing signal counts as unknown, never safe. "
             "False positives hurt people, so we only block at very high confidence and keep challenges under five percent. "
             "And behaviour data is personal data, so we keep derived features only. "
             "Our targets: stop at least 80 percent of scam-type fraud while challenging at most 5 percent of honest payments, let 90 percent of honest big transfers through with no delay, and decide in under 50 milliseconds.")


def main(src, out):
    prs = Presentation(src)
    sl = list(prs.slides)
    for s in sl:
        strip_instructions(s)
    slide1(sl[0]); slide2(sl[1]); slide3(sl[2]); slide4(sl[3]); slide5(sl[4])
    for s in sl:  # keep "SLIDE X OF 5" on the canvas
        for sh in s.shapes:
            if sh.name == "Text 1":
                sh.left = Inches(11.3); sh.width = Inches(1.55)
                sh.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    prs.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
